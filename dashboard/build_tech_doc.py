from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "疯狂原始人_机器人系统技术说明.docx"
BLUE = "2E74B5"; NAVY = "20364A"; ORANGE = "E7682B"; LIGHT = "E8EEF5"; CREAM = "FFF4E6"; GRAY = "667085"

def font(run, size=11, bold=False, color=None, name="Noto Sans CJK SC"):
    run.font.name = name
    rf=run._element.get_or_add_rPr().rFonts
    for key in ("ascii","hAnsi","eastAsia","cs"): rf.set(qn("w:"+key),name)
    run.font.size = Pt(size); run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar = OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for edge, value in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = tcMar.find(qn("w:"+edge))
        if node is None: node=OxmlElement("w:"+edge); tcMar.append(node)
        node.set(qn("w:w"),str(value)); node.set(qn("w:type"),"dxa")

def set_table_widths(table, widths):
    table.autofit = False
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    tblPr=table._tbl.tblPr; tblW=tblPr.first_child_found_in("w:tblW")
    if tblW is None: tblW=OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"),str(sum(widths))); tblW.set(qn("w:type"),"dxa")
    tblInd=tblPr.first_child_found_in("w:tblInd")
    if tblInd is None: tblInd=OxmlElement("w:tblInd"); tblPr.append(tblInd)
    tblInd.set(qn("w:w"),"120"); tblInd.set(qn("w:type"),"dxa")
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcW=cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tcW is None: tcW=OxmlElement("w:tcW"); cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"),str(widths[i])); tcW.set(qn("w:type"),"dxa"); cell_margins(cell); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")

def add_bullets(doc, items):
    for item in items:
        p=doc.add_paragraph(style="List Bullet"); p.add_run(item)

def add_caption(doc, text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(9)
    font(p.add_run(text),9,False,GRAY)

def add_kv_table(doc, rows):
    t=doc.add_table(rows=1, cols=3); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.style="Table Grid"
    hdr=t.rows[0].cells
    for i,x in enumerate(("层级","技术 / 模型","作用")): hdr[i].text=x; shade(hdr[i],LIGHT)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row): cells[i].text=x
    set_table_widths(t,[1750,2850,4760])
    for row_i,row in enumerate(t.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs: font(r,9.5,row_i==0,NAVY if row_i==0 else None)
    return t

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
styles=doc.styles
normal=styles["Normal"]; normal.font.name="Noto Sans CJK SC"; normal._element.rPr.rFonts.set(qn("w:eastAsia"),"Noto Sans CJK SC"); normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.25
for level,size,before,after,color in ((1,16,18,10,BLUE),(2,13,14,7,BLUE),(3,12,10,5,"1F4D78")):
    s=styles[f"Heading {level}"]; s.font.name="Noto Sans CJK SC"; s._element.rPr.rFonts.set(qn("w:eastAsia"),"Noto Sans CJK SC"); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
for style_name in ("List Bullet","List Number"):
    s=styles[style_name]; s.font.name="Noto Sans CJK SC"; s._element.rPr.rFonts.set(qn("w:eastAsia"),"Noto Sans CJK SC"); s.font.size=Pt(11); s.paragraph_format.left_indent=Inches(.375); s.paragraph_format.first_line_indent=Inches(-.188); s.paragraph_format.space_after=Pt(4); s.paragraph_format.line_spacing=1.25

header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(header.add_run("疯狂原始人 · 技术说明"),9,True,GRAY)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(footer.add_run("A1R 视觉机械臂原型系统  |  2026"),9,False,GRAY)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(54); p.paragraph_format.space_after=Pt(10); font(p.add_run("CRAZY CAVEMAN ROBOTICS"),10,True,ORANGE)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(8); font(p.add_run("疯狂原始人"),30,True,NAVY)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(20); font(p.add_run("A1R 视觉机械臂系统技术说明"),17,False,BLUE)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(22); font(p.add_run("从边缘视觉检测、目标筛选、闭环靠近，到示教动作库与可视化任务中心的完整实现。"),12,False,GRAY)

call=doc.add_table(rows=1,cols=1); call.style="Table Grid"; set_table_widths(call,[9360]); shade(call.cell(0,0),CREAM)
cp=call.cell(0,0).paragraphs[0]; font(cp.add_run("当前落地能力  "),11,True,ORANGE); font(cp.add_run("reCamera 实时检测地面物品，A1R 根据目标位置分步平滑靠近；同时支持零力示教、动作录制、分类与回放。"),11,False,NAVY)

doc.add_paragraph()
doc.add_picture(str(ROOT/"crazy_caveman_ui.png"),width=Inches(6.2)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break(); add_heading(doc,"1. 系统概览",1)
doc.add_paragraph("系统采用“边缘视觉 + 任务控制 + 机械臂执行”的分层结构。reCamera 在相机端完成实时推理，机械臂板接收检测结果并生成短步长末端位移，浏览器任务中心用于启动任务、观察状态、管理示教动作及执行急停。")
add_kv_table(doc,[
    ("感知层","Seeed Studio reCamera 2002（64GB）","采集 1080p 画面，运行边缘目标检测并输出检测框。"),
    ("检测模型","YOLO11n / SSCMA 推理节点","识别 COCO 类别，输出标签、置信度、中心坐标与框尺寸。"),
    ("视觉编排","Node-RED + HTTP JSON","缓存最新推理结果，并提供 /api/latest-detection。"),
    ("任务层","Python 3 + ThreadingHTTPServer","目标筛选、稳定识别、闭环跟随、运行状态与启停接口。"),
    ("控制层","C++17 + OneroArm SDK","串口/CAN 控制、MoveP 逆解、轨迹分段、示教录制与回放。"),
    ("交互层","HTML/CSS/JavaScript + Three.js","任务化界面、动作分类、3D 末端控制和七关节检查。"),
])

add_heading(doc,"2. 数据与控制链路",1)
for step in [
    "相机采集画面，YOLO11n 在 reCamera 本地推理。",
    "Node-RED 将最新 labels、boxes、resolution 结果暴露为 HTTP JSON。",
    "视觉靠近服务过滤 person、toilet、家具等干扰，并优先选择小型可拾取物体。",
    "同一目标连续稳定识别 3 帧后，根据图像中心误差计算横向 Y 位移，并依据框尺寸计算下降 Z 位移。",
    "每次合成位移限制在约 4.5 cm，调用 A1R /api/end-move，以最高 0.20 的速度比例分段执行。",
    "目标已居中且框足够大时停止下降并保持跟随；用户可随时停止任务或急停。",
]:
    p=doc.add_paragraph(style="List Number"); p.add_run(step)

add_heading(doc,"3. 视觉模型与目标策略",1)
doc.add_picture(str(ROOT/"recamera_preview.png"),width=Inches(6.45)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
add_caption(doc,"图 2  reCamera 实时预览与边缘检测界面")
add_heading(doc,"3.1 当前部署模型：YOLO11n",2)
doc.add_paragraph("YOLO11n 是当前实时检测主模型，特点是参数量小、推理延迟低，适合在 reCamera 边缘设备上持续运行。系统直接使用检测框，不依赖云端请求，因此相机到控制器链路可保持低成本和可调试性。")
add_bullets(doc,[
    "输入分辨率：检测结果按 1280 × 720 坐标系输出；相机源为 1080p。",
    "主要输出：类别标签、置信度、框中心 x/y、框宽高及推理性能信息。",
    "过滤策略：忽略人员，同时排除机械臂常见的大型误检框和固定家具类别。",
    "候选优先级：书本、球、瓶子、杯子、遥控器、手机、胡萝卜等小型可拾取类别优先。",
])
add_heading(doc,"3.2 大模型与 VLA 的定位",2)
doc.add_paragraph("系统已经预留视觉大模型 API，可在后续用于开放词汇识别、复杂场景语义理解与任务规划；π0.5 等 VLA 模型也可作为端到端动作策略的研究方向。但当前“点击捡东西→机械臂靠近”的闭环未使用 π0.5，也不依赖云端大模型，实际控制由 YOLO11n 检测和确定性运动规则完成。")

add_heading(doc,"4. 机械臂控制与安全约束",1)
add_bullets(doc,[
    "位置控制：使用 OneroArm SDK 的 MoveP 进行末端笛卡尔运动与逆运动学求解。",
    "平滑性：较长位移按约 2 cm 分段并使用缓冲轨迹连续执行。",
    "单次约束：手动接口最大 8 cm；视觉闭环采用更小的约 4.5 cm 步长。",
    "模式互斥：零力示教、录制/回放、位置控制和视觉靠近不会同时抢占机械臂。",
    "急停：任务面板的停止按钮同时关闭视觉任务并请求取消当前轨迹。",
    "设备稳定性：USB-CAN 使用 /dev/serial/by-id 稳定标识，避免 ttyACM0/1 变化。",
])

doc.add_page_break(); add_heading(doc,"5. 可视化任务中心",1)
doc.add_paragraph("界面品牌为“疯狂原始人”，按用户任务而不是底层接口组织功能。主页面优先展示视觉拾取，其次是动作录制与任务库，高级三维/关节控制放在下方。")
add_kv_table(doc,[
    ("视觉拾取","捡东西 · 开始靠近 / 停止靠近","启动持续检测与跟随，不控制夹爪。"),
    ("复位姿态","复位、坐下、站立、归位类动作","用于基础姿态和任务起点。"),
    ("互动表演","挥手、握手、摇、打招呼类动作","用于展示和人机互动。"),
    ("作业任务","拿东西、拿伞、拿行李等动作","用于真实操作任务的录制与回放。"),
])
add_heading(doc,"6. 服务与接口",1)
add_kv_table(doc,[
    ("A1R 控制服务","端口 8080","/api/status、/api/end-move、/api/teaching/enable、录制/回放/停止。"),
    ("视觉靠近服务","端口 8090","/api/approach/start、/stop、/status。"),
    ("reCamera Node-RED","端口 1880","/dashboard/preview、/api/latest-detection。"),
])
add_heading(doc,"7. 当前边界与下一步",1)
add_bullets(doc,[
    "当前为单目视觉，深度主要通过目标框尺寸近似；精确抓取需增加深度或相机—机械臂标定。",
    "当前只完成“靠近”，尚未接入林巧手夹爪闭合、抓取确认和失败恢复。",
    "建议增加目标锁定 ID、卡尔曼滤波和坐标标定，提高快速移动时的跟随稳定性。",
    "可将大模型用于目标名称理解与任务分解，但实时运动闭环应继续保留本地限幅和急停。",
])

doc.save(OUT)
print(OUT)
